import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt
import io
import os

from pptx import Presentation
from pptx.util import Inches, Pt

st.set_page_config(page_title="Consolidada IMB", layout="wide")

# =========================
# ESTILO
# =========================
st.markdown("""
<style>
html, body, [class*="css"] {
    font-family: 'Noto Sans', sans-serif;
}

.titulo {
    color:#9F2241;
    font-weight:800;
    text-align:center;
}

thead tr th {
    background-color:#235B4E !important;
    color:white !important;
    font-weight:bold !important;
}

[data-testid="stDataFrame"] {
    border:2px solid #235B4E;
}
</style>
""", unsafe_allow_html=True)

st.markdown(
    "<h1 class='titulo'>📊 Dashboard IMSS BIENESTAR - Abasto</h1>",
    unsafe_allow_html=True
)

# =========================
# ARCHIVO
# =========================
# =========================
# CARGA DE ARCHIVO
# =========================

st.sidebar.markdown("## 📂 Cargar archivo")

archivo = st.sidebar.file_uploader(
    "Subir Consolidada",
    type=["xlsb", "xlsx"]
)

archivo_inv = st.sidebar.file_uploader(
    "Subir Inventario",
    type=["xlsx", "xlsb"]
)

archivo_cpm = st.sidebar.file_uploader(
    "Subir CPM",
    type=["xlsx", "xlsb"]
)
if (
    archivo is None or
    archivo_inv is None or
    archivo_cpm is None
):
    st.warning("⬅ Debes cargar todos los archivos.")
    st.stop()
@st.cache_data(show_spinner=False)
def cargar():
    reemplazos = {
        "MICHOACÁN": "MICHOACAN",
        "MICHOACÁN DE OCAMPO": "MICHOACAN",
        "MICHOACAN DE OCAMPO": "MICHOACAN",
        "BAJA SUR": "BAJA CALIFORNIA SUR",
        "B.C.S.": "BAJA CALIFORNIA SUR",
        "BCS": "BAJA CALIFORNIA SUR"
    }
    df = pd.read_excel(
        archivo,
        engine="pyxlsb",
        usecols=[0, 3, 17, 18, 24, 27, 55]
    )
    df.columns = df.columns.str.strip()
    # =========================
    # INVENTARIO
    # =========================

    inv = pd.read_excel(
        archivo_inv
    )

    inv.columns = inv.columns.str.strip()

    col_inv_estado = inv.columns[0]
    col_inv_clave = inv.columns[3]
    col_inv_piezas = inv.columns[5]

    inv["ENTIDAD"] = (
        inv[col_inv_estado]
        .astype(str)
        .str.strip()
        .str.upper()
        .str.normalize("NFKD")
        .str.encode("ascii", errors="ignore")
        .str.decode("utf-8")
    )

    inv["ENTIDAD"] = (
        inv["ENTIDAD"]
        .replace(reemplazos)
    )

    inv["CLAVE"] = (
        inv[col_inv_clave]
        .astype(str)
        .str.strip()
    )

    inv["CLAVE_MERGE"] = (
        inv["CLAVE"]
        .str.replace(".", "", regex=False)
        .str.replace(" ", "", regex=False)
    )

    inv["PIEZAS_INV"] = (
        inv[col_inv_piezas]
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.strip()
    )
    # =========================
    # FILTRAR INVENTARIO VIGENTE
    # =========================

    col_inv_estatus = inv.columns[7]

    inv["ESTATUS"] = (
        inv[col_inv_estatus]
        .astype(str)
        .str.strip()
        .str.upper()
    )

    inv = inv[
        ~inv["ESTATUS"].str.contains(
            "CADUC",
            na=False
        )
    ]
    inv["PIEZAS_INV"] = pd.to_numeric(
        inv["PIEZAS_INV"],
        errors="coerce"
    ).fillna(0)

    # CONSOLIDAR INVENTARIO
    inv_group = (
        inv.groupby(
            ["ENTIDAD", "CLAVE_MERGE"],
            dropna=False
        )["PIEZAS_INV"]
        .sum()
        .reset_index()
    )
    # =========================
    # CPM
    # =========================

    cpm = pd.read_excel(
        archivo_cpm
    )

    cpm.columns = cpm.columns.str.strip()

    col_cpm_estado = cpm.columns[0]
    col_cpm_clave = cpm.columns[1]
    col_cpm = cpm.columns[3]

    cpm["ENTIDAD"] = (
        cpm[col_cpm_estado]
        .astype(str)
        .str.strip()
        .str.upper()
        .str.normalize("NFKD")
        .str.encode("ascii", errors="ignore")
        .str.decode("utf-8")
    )

    cpm["ENTIDAD"] = (
        cpm["ENTIDAD"]
        .replace(reemplazos)
    )

    cpm["CLAVE"] = (
        cpm[col_cpm_clave]
        .astype(str)
        .str.strip()
    )

    cpm["CLAVE_MERGE"] = (
        cpm["CLAVE"]
        .str.replace(".", "", regex=False)
        .str.replace(" ", "", regex=False)
    )

    cpm["CPM"] = (
        cpm[col_cpm]
        .astype(str)
        .str.replace(",", "", regex=False)
        .str.strip()
    )

    cpm["CPM"] = pd.to_numeric(
        cpm["CPM"],
        errors="coerce"
    ).fillna(0)

    cpm_group = (
        cpm.groupby(
            ["ENTIDAD", "CLAVE_MERGE"],
            dropna=False
        )["CPM"]
        .sum()
        .reset_index()
    )
    # =========================
        # =========================
    # LIMPIAR CONSOLIDADA
    # =========================

    df["ENTIDAD"] = (
        df[df.columns[1]]
        .astype(str)
        .str.strip()
        .str.upper()
        .str.normalize("NFKD")
        .str.encode("ascii", errors="ignore")
        .str.decode("utf-8")
    )

    df["ENTIDAD"] = (
        df["ENTIDAD"]
        .replace(reemplazos)
    )

    df["CLAVE"] = (
        df[df.columns[2]]
        .astype(str)
        .str.strip()
    )

    df["CLAVE_MERGE"] = (
        df["CLAVE"]
        .str.replace(".", "", regex=False)
        .str.replace(" ", "", regex=False)
    )

    # =========================
    # MERGE INVENTARIO
    # =========================

    df = df.merge(
        inv_group,
        how="left",
        on=["ENTIDAD", "CLAVE_MERGE"],
    )

    df["PIEZAS_INV"] = (
        df["PIEZAS_INV"]
        .fillna(0)
    )

    # =========================
    # MERGE CPM
    # =========================

    df = df.merge(
        cpm_group,
        how="left",
        on=["ENTIDAD", "CLAVE_MERGE"],
    )

    df["CPM"] = (
        df["CPM"]
        .fillna(0)
    )

    # =========================
    # AGREGAR CLAVES SIN EMISIÓN
    # =========================

    base_extra = pd.merge(
        inv_group,
        cpm_group,
        how="outer",
        on=["ENTIDAD", "CLAVE_MERGE"]
    )

    existentes = set(
        zip(df["ENTIDAD"], df["CLAVE_MERGE"])
    )

    base_extra = base_extra[
        ~base_extra.apply(
            lambda x: (
                x["ENTIDAD"],
                x["CLAVE_MERGE"]
            ) in existentes,
            axis=1
        )
    ]

    if len(base_extra) > 0:

        extra = pd.DataFrame({
            df.columns[0]: "SIN EMISION",
            df.columns[1]: base_extra["ENTIDAD"],
            df.columns[2]: base_extra["CLAVE_MERGE"],	
            df.columns[3]: "SIN DESCRIPCIÓN",
            df.columns[4]: 0,
            df.columns[5]: 0,
            df.columns[6]: 0,
            "PIEZAS_INV": base_extra["PIEZAS_INV"],
            "CPM": base_extra["CPM"]
        })

        df = pd.concat(
            [df, extra],
            ignore_index=True
        )
    return df, inv_group, cpm_group

df, inv_group, cpm_group = cargar()
# =========================
# COLUMNAS
# =========================
col_orden = df.columns[0]
col_estado = df.columns[1]
col_clave = df.columns[2]
col_descripcion = df.columns[3]
col_precio = df.columns[4]
col_emitidas = df.columns[5]
col_entregadas = df.columns[6]
# =========================
# BASE BUSCADOR
# =========================

df["PIEZAS_INV"] = pd.to_numeric(
    df["PIEZAS_INV"],
    errors="coerce"
).fillna(0)

df["CPM"] = pd.to_numeric(
    df["CPM"],
    errors="coerce"
).fillna(0)
base_busqueda = pd.merge(
    inv_group,
    cpm_group,
    how="outer",
    on=["ENTIDAD", "CLAVE_MERGE"]
)

base_busqueda["PIEZAS_INV"] = (
    base_busqueda["PIEZAS_INV"]
    .fillna(0)
)

base_busqueda["CPM"] = (
    base_busqueda["CPM"]
    .fillna(0)
)

emitido = (
    df.groupby(
        ["ENTIDAD", "CLAVE_MERGE"],
        dropna=False
    )[col_emitidas]
    .sum()
    .reset_index()
)

entregado = (
    df.groupby(
        ["ENTIDAD", "CLAVE_MERGE"],
        dropna=False
    )[col_entregadas]
    .sum()
    .reset_index()
)

descripcion = (
    df.groupby(
        ["ENTIDAD", "CLAVE_MERGE"],
        dropna=False
    )[col_descripcion]
    .first()
    .reset_index()
)

base_busqueda = base_busqueda.merge(
    emitido,
    how="left",
    on=["ENTIDAD", "CLAVE_MERGE"]
)

base_busqueda = base_busqueda.merge(
    entregado,
    how="left",
    on=["ENTIDAD", "CLAVE_MERGE"]
)

base_busqueda = base_busqueda.merge(
    descripcion,
    how="left",
    on=["ENTIDAD", "CLAVE_MERGE"]
)
base_busqueda = base_busqueda.merge(
    df[
        ["CLAVE_MERGE", "CLAVE"]
    ].drop_duplicates(),
    how="left",
    on="CLAVE_MERGE"
)
base_busqueda[col_emitidas] = (
    base_busqueda[col_emitidas]
    .fillna(0)
)

base_busqueda[col_entregadas] = (
    base_busqueda[col_entregadas]
    .fillna(0)
)


# =========================
# FORMATOS
# =========================
def fmt(x):
    return f"{int(x):,}"

def fmt_money(x):
    return f"${int(x):,}"

# =========================
# MÉTRICAS
# =========================
def calcular_metricas(df_estado):

    df_estado = df_estado.copy()
    # =========================
    # ABASTO
    # =========================

    df_estado["PIEZAS_INV"] = pd.to_numeric(
        df_estado["PIEZAS_INV"],
        errors="coerce"
    ).fillna(0)

    df_estado["CPM"] = pd.to_numeric(
        df_estado["CPM"],
        errors="coerce"
    ).fillna(0)

    df_estado["NIVEL_ABASTO"] = 0.0

    mask = df_estado["CPM"] > 0

    df_estado.loc[mask, "NIVEL_ABASTO"] = (
        df_estado.loc[mask, "PIEZAS_INV"] /
        df_estado.loc[mask, "CPM"]
    )

    # CLAVES CON INVENTARIO Y CPM 0
    df_estado.loc[
        (df_estado["PIEZAS_INV"] > 0) &
        (df_estado["CPM"] == 0),
        "NIVEL_ABASTO"
    ] = 2
    # =========================
    # CLASIFICACIÓN
    # =========================

    def clasificar(x):

        if x == 0:
            return "Agotado"

        elif x < 1:
            return "Próx agotarse"

        elif x <= 1.5:
            return "Bajo"

        elif x <= 5:
            return "Óptimo"

        else:
            return "Sobre stock"

    df_estado["CLASIFICACION"] = (
        df_estado["NIVEL_ABASTO"]
        .apply(clasificar)
    )

    df_estado["Emitido"] = pd.to_numeric(
        df_estado[col_emitidas],
        errors="coerce"
    ).fillna(0)

    df_estado["Entregado"] = pd.to_numeric(
        df_estado[col_entregadas],
        errors="coerce"
    ).fillna(0)

    df_estado[col_precio] = pd.to_numeric(
        df_estado[col_precio],
        errors="coerce"
    ).fillna(0)
    df_estado[col_precio] = pd.to_numeric(
        df_estado[col_precio],
        errors="coerce"
    ).fillna(0)

    df_estado["Transito"] = (
        df_estado["Emitido"] - df_estado["Entregado"]
    ).clip(lower=0)

    df_estado["Monto_Emitido"] = df_estado["Emitido"] * df_estado[col_precio]
    df_estado["Monto_Entregado"] = df_estado["Entregado"] * df_estado[col_precio]
    df_estado["Monto_Transito"] = df_estado["Monto_Emitido"] - df_estado["Monto_Entregado"]

    p = df_estado["Emitido"].sum()
    e = df_estado["Entregado"].sum()
    t = df_estado["Transito"].sum()

    m_p = df_estado["Monto_Emitido"].sum()
    m_e = df_estado["Monto_Entregado"].sum()
    m_t = df_estado["Monto_Transito"].sum()

    o_total = df_estado[col_orden].nunique()
    c_total = df_estado[col_clave].nunique()

    o_ent = df_estado[df_estado["Entregado"] >= df_estado["Emitido"]][col_orden].nunique()
    o_tran = df_estado[df_estado["Entregado"] < df_estado["Emitido"]][col_orden].nunique()

    c_ent = df_estado[df_estado["Entregado"] >= df_estado["Emitido"]][col_clave].nunique()
    c_tran = df_estado[df_estado["Entregado"] < df_estado["Emitido"]][col_clave].nunique()

    inv_total = (
    df_estado[
        ["ENTIDAD", "CLAVE", "PIEZAS_INV"]
    ]
    .drop_duplicates()
    ["PIEZAS_INV"]
    .sum()
)
    cpm_total = (
    df_estado[
        ["ENTIDAD", "CLAVE", "CPM"]
    ]
    .drop_duplicates()
    ["CPM"]
    .sum()
)

    agotado = (
        df_estado["CLASIFICACION"]
        == "Agotado"
    ).sum()

    proximo = (
        df_estado["CLASIFICACION"]
        == "Próx agotarse"
    ).sum()

    bajo = (
        df_estado["CLASIFICACION"]
        == "Bajo"
    ).sum()

    optimo = (
        df_estado["CLASIFICACION"]
        == "Óptimo"
    ).sum()

    sobre = (
        df_estado["CLASIFICACION"]
        == "Sobre stock"
    ).sum()

    nivel_promedio = (
        df_estado["NIVEL_ABASTO"]
        .mean()
    )    
    tabla = pd.DataFrame([
        ["Piezas", fmt(p), fmt(t), fmt(e)],
        ["Claves", fmt(c_total), fmt(c_tran), fmt(c_ent)],
        ["Órdenes", fmt(o_total), fmt(o_tran), fmt(o_ent)],
        ["Montos", fmt_money(m_p), fmt_money(m_t), fmt_money(m_e)]
    ], columns=["Métrica", "Emitido", "En tránsito", "Entregado"])

    return {
        "inv_total": inv_total,
        "cpm_total": cpm_total,
        "agotado": agotado,
        "proximo": proximo,
        "bajo": bajo,
        "optimo": optimo,
        "sobre": sobre,
        "nivel_promedio": nivel_promedio,
        "p": p,
        "e": e,
        "t": t,
        "m_p": m_p,
        "m_e": m_e,
        "m_t": m_t,
        "o_total": o_total,
        "o_ent": o_ent,
        "o_tran": o_tran,
        "c_total": c_total,
        "c_ent": c_ent,
        "c_tran": c_tran,
        "tabla": tabla
    }

# =========================
# GRÁFICAS
# =========================
def grafica(titulo, valores, labels, colores):

    fig, ax = plt.subplots(figsize=(10.5, 3.6))

    ax.barh(labels, valores, color=colores)

    ax.set_title(
        titulo,
        fontsize=15,
        fontweight="bold"
    )

    ax.set_xticks([])
    ax.tick_params(axis="y", labelsize=11)

    max_val = max(valores) if max(valores) > 0 else 1

    for i, v in enumerate(valores):
        ax.text(
            v + (max_val * 0.01),
            i,
            f"{int(v):,}",
            va="center",
            fontsize=11,
            fontweight="bold"
        )

    ax.set_xlim(0, max_val * 1.18)
    ax.set_facecolor("none")
    fig.patch.set_alpha(0)

    plt.tight_layout()

    return fig

def crear_graficas(m):

    fig_piezas = grafica(
        "Piezas",
        [m["p"], m["e"], m["t"]],
        ["Emitido", "Entregado", "Tránsito"],
        ["#235B4E", "#9F2241", "#B38E5D"]
    )

    fig_montos = grafica(
        "Montos ($)",
        [m["m_p"], m["m_e"], m["m_t"]],
        ["Emitido", "Entregado", "Tránsito"],
        ["#235B4E", "#9F2241", "#B38E5D"]
    )

    fig_ordenes = grafica(
        "Órdenes",
        [m["o_total"], m["o_ent"], m["o_tran"]],
        ["Total", "Entregadas", "En tránsito"],
        ["#235B4E", "#9F2241", "#B38E5D"]
    )

    fig_claves = grafica(
        "Claves",
        [m["c_total"], m["c_ent"], m["c_tran"]],
        ["Total", "Entregadas", "En tránsito"],
        ["#235B4E", "#9F2241", "#B38E5D"]
    )

    return fig_piezas, fig_montos, fig_ordenes, fig_claves

# =========================
# RESUMEN GENERAL NACIONAL
# =========================
metricas_general = calcular_metricas(df)

st.markdown("## 🌎 Resumen general nacional")

g1, g2, g3 = st.columns(3)
g4, g5, g6 = st.columns(3)

with g4:
    st.metric(
        "📦 Inventario",
        fmt(metricas_general["inv_total"])
    )

with g5:
    st.metric(
        "📈 CPM",
        fmt(metricas_general["cpm_total"])
    )

with g1:
    st.metric(
        "🔵 Total emitido",
        fmt(metricas_general["p"])
    )

with g2:
    st.metric(
        "🟢 Total entregado",
        fmt(metricas_general["e"])
    )

with g3:
    st.metric(
        "🟡 Total en tránsito",
        fmt(metricas_general["t"])
    )

st.markdown("### 📋 Tabla nacional")

col_n1, col_n2, col_n3 = st.columns([1, 3, 1])

with col_n2:
    st.dataframe(
        metricas_general["tabla"],
        use_container_width=True,
        hide_index=True
    )

st.divider()
# =========================
# BUSCADOR DE CLAVE
# =========================

st.markdown("## 🔎 Buscar clave nacional")

clave_busqueda = st.text_input(
    "Ingresar clave"
)
clave_busqueda = (
    clave_busqueda
    .replace(".", "")
    .replace(" ", "")
    .strip()
)

if clave_busqueda:

    df_busqueda = base_busqueda.copy()

    df_busqueda = df_busqueda[
        df_busqueda["CLAVE"]
        .astype(str)
        .str.contains(
            clave_busqueda,
            case=False,
            na=False
        )
    ]

    if len(df_busqueda) > 0:

        df_busqueda["NIVEL_ABASTO"] = 0.0

        mask = df_busqueda["CPM"] > 0

        df_busqueda.loc[mask, "NIVEL_ABASTO"] = (
            df_busqueda.loc[mask, "PIEZAS_INV"] /
            df_busqueda.loc[mask, "CPM"]
        )

        df_busqueda.loc[
            (df_busqueda["PIEZAS_INV"] > 0) &
            (df_busqueda["CPM"] == 0),
            "NIVEL_ABASTO"
        ] = 2

        def clasificar_busqueda(x):

            if x == 0:
                return "Agotado"

            elif x < 1:
                return "Próx agotarse"

            elif x <= 1.5:
                return "Bajo"

            elif x <= 5:
                return "Óptimo"

            else:
                return "Sobre stock"

        df_busqueda["Clasificación"] = (
            df_busqueda["NIVEL_ABASTO"]
            .apply(clasificar_busqueda)
        )

        tabla_busqueda = df_busqueda.copy()
        tabla_busqueda["Tránsito"] = (
            tabla_busqueda[col_emitidas] -
            tabla_busqueda[col_entregadas]
        ).clip(lower=0)

        tabla_busqueda["Tránsito"] = (
            tabla_busqueda[col_emitidas] -
            tabla_busqueda[col_entregadas]
        ).clip(lower=0)

        tabla_busqueda = tabla_busqueda.rename(columns={
            "ENTIDAD": "Estado",
            "CLAVE": "Clave",
            "PIEZAS_INV": "Inventario",
            "CPM": "CPM",
            "NIVEL_ABASTO": "Nivel",
            "Clasificación": "Clasificación",
            col_entregadas: "Entregado",
            col_emitidas: "Emitido"
        })

        tabla_busqueda["Nivel"] = (
            tabla_busqueda["Nivel"]
            .round(2)
        )

        descripcion = (
            str(df_busqueda.iloc[0][df.columns[3]])
        )

        st.markdown(
            f"### 🔎 Clave encontrada: {descripcion}"
        )

        tabla_busqueda = tabla_busqueda.sort_values(
            "Inventario",
            ascending=False
        )

        # QUITAR COLUMNAS TÉCNICAS
        tabla_busqueda = tabla_busqueda.drop(
            columns=["CLAVE_MERGE"],
            errors="ignore"
        )

        st.dataframe(
            tabla_busqueda,
            use_container_width=True,
            hide_index=True
        )
# =========================
# FILTRO
# =========================
estado_sel = st.selectbox(
    "📍 Estado",
    sorted(df[col_estado].dropna().unique())
)

df_f = df[df[col_estado] == estado_sel].copy()
# =========================
# RESUMEN ABASTO
# =========================

df_f = df_f.copy()

df_f["NIVEL_ABASTO"] = 0.0

mask = df_f["CPM"] > 0

df_f.loc[mask, "NIVEL_ABASTO"] = (
    df_f.loc[mask, "PIEZAS_INV"] /
    df_f.loc[mask, "CPM"]
)

df_f.loc[
    (df_f["PIEZAS_INV"] > 0) &
    (df_f["CPM"] == 0),
    "NIVEL_ABASTO"
] = 2

def clasificar_tabla(x):

    if x == 0:
        return "Agotado"

    elif x < 1:
        return "Próx agotarse"

    elif x <= 1.5:
        return "Bajo"

    elif x <= 5:
        return "Óptimo"

    else:
        return "Sobre stock"

df_f["CLASIFICACION"] = (
    df_f["NIVEL_ABASTO"]
    .apply(clasificar_tabla)
)
abasto_claves = (
    df_f.groupby(
        [col_clave],
        dropna=False
    )
    .agg({
        "PIEZAS_INV": "sum",
        "CPM": "sum",
        "NIVEL_ABASTO": "first",
        "CLASIFICACION": "first"
    })
    .reset_index()
)

abasto_claves.columns = [
    "Clave",
    "Inventario",
    "CPM",
    "Nivel",
    "Clasificación"
]

abasto_claves["Nivel"] = (
    abasto_claves["Nivel"]
    .round(2)
)

abasto_claves = abasto_claves.sort_values(
    "Nivel"
)

metricas = calcular_metricas(df_f)

tabla_base = metricas["tabla"]

fig_piezas, fig_montos, fig_ordenes, fig_claves = crear_graficas(metricas)

# =========================
# TABLA STREAMLIT
# =========================
st.markdown(f"## 📋 Tabla operativa - {estado_sel}")

col1, col2, col3 = st.columns([1, 3, 1])

with col2:
    st.dataframe(
        tabla_base,
        use_container_width=True,
        hide_index=True
    )

# =========================
# VISUALIZACIÓN
# =========================
resumen_abasto = (
    abasto_claves["Clasificación"]
    .value_counts()
    .reset_index()
)

resumen_abasto.columns = [
    "Clasificación",
    "Total claves"
]

st.markdown(f"## 🏥 Resumen de abasto - {estado_sel}")

st.dataframe(
    resumen_abasto,
    use_container_width=True,
    hide_index=True
)
st.markdown(f"## 📊 Visualización - {estado_sel}")

c1, c2 = st.columns(2)
c3, c4 = st.columns(2)

with c1:
    st.pyplot(fig_piezas)

with c2:
    st.pyplot(fig_montos)

with c3:
    st.pyplot(fig_ordenes)

with c4:
    st.pyplot(fig_claves)

# =========================
# POWERPOINT
# =========================
MACHOTE = "MACHOTE_PRESENTACION.pptx"

def obtener_layout(prs):
    if len(prs.slide_layouts) > 5:
        return prs.slide_layouts[5]
    elif len(prs.slide_layouts) > 0:
        return prs.slide_layouts[0]
    else:
        raise Exception("El PowerPoint no tiene layouts disponibles.")

def limpiar_slide(slide):
    for shape in list(slide.shapes):
        try:
            sp = shape._element
            sp.getparent().remove(sp)
        except Exception:
            pass

def agregar_contenido_slide(
    slide,
    tabla,
    fig_piezas,
    fig_montos,
    fig_ordenes,
    fig_claves,
    estado
):

    txBox = slide.shapes.add_textbox(
        Inches(0.6),
        Inches(0.2),
        Inches(8),
        Inches(0.5)
    )

    tf = txBox.text_frame
    tf.text = f"Abasto - {str(estado).title()}"

    p_title = tf.paragraphs[0]
    p_title.font.size = Pt(30)
    p_title.font.bold = True

    fig_tab, ax = plt.subplots(figsize=(10, 2))
    ax.axis("off")

    table = ax.table(
        cellText=tabla.values,
        colLabels=tabla.columns,
        loc="center",
        cellLoc="center"
    )

    table.scale(1.4, 1.5)

    for (r, c), cell in table.get_celld().items():
        cell.set_edgecolor("black")

        if r == 0:
            cell.set_facecolor("#235B4E")
            cell.set_text_props(
                color="white",
                weight="bold",
                size=10
            )
        else:
            cell.set_facecolor((1, 1, 1, 0))
            cell.set_text_props(size=10)

    fig_tab.patch.set_alpha(0)

    img_tab = io.BytesIO()

    plt.savefig(
        img_tab,
        format="png",
        bbox_inches="tight",
        transparent=True
    )

    img_tab.seek(0)

    plt.close(fig_tab)

    slide.shapes.add_picture(
        img_tab,
        Inches(1.75),
        Inches(0.75),
        width=Inches(9.7)
    )

    def add(fig, x, y):

        img = io.BytesIO()

        fig.savefig(
            img,
            format="png",
            bbox_inches="tight",
            transparent=True
        )

        img.seek(0)

        slide.shapes.add_picture(
            img,
            Inches(x),
            Inches(y),
            width=Inches(4.9)
        )

        plt.close(fig)

    add(fig_piezas, 0.55, 2.05)
    add(fig_montos, 6.55, 2.05)
    add(fig_ordenes, 0.55, 4.25)
    add(fig_claves, 6.55, 4.25)

def exportar_ppt_estado(
    tabla_base,
    fig_piezas,
    fig_montos,
    fig_ordenes,
    fig_claves,
    estado_sel
):

    prs = Presentation(MACHOTE) if os.path.exists(MACHOTE) else Presentation()

    if len(prs.slides) == 0:
        slide = prs.slides.add_slide(obtener_layout(prs))
    else:
        slide = prs.slides[0]

    limpiar_slide(slide)

    agregar_contenido_slide(
        slide,
        tabla_base,
        fig_piezas,
        fig_montos,
        fig_ordenes,
        fig_claves,
        estado_sel
    )

    salida = f"Dashboard_IMB_{estado_sel}.pptx"

    prs.save(salida)

    return salida

def exportar_ppt_todos_estados():

    prs = Presentation(MACHOTE) if os.path.exists(MACHOTE) else Presentation()

    while len(prs.slides._sldIdLst) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]

    estados = sorted(df[col_estado].dropna().unique())

    barra = st.progress(0)
    total = len(estados)

    for i, estado in enumerate(estados):

        df_estado = df[df[col_estado] == estado].copy()

        m = calcular_metricas(df_estado)

        fig_p, fig_m, fig_o, fig_c = crear_graficas(m)

        slide = prs.slides.add_slide(obtener_layout(prs))

        limpiar_slide(slide)

        agregar_contenido_slide(
            slide,
            m["tabla"],
            fig_p,
            fig_m,
            fig_o,
            fig_c,
            estado
        )

        barra.progress((i + 1) / total)

    salida = "Dashboard_Nacional_IMB_Todos_Estados.pptx"

    prs.save(salida)

    return salida

# =========================
# BOTONES
# =========================
st.markdown("## 📤 Exportación PPT")

col_exp1, col_exp2 = st.columns(2)

with col_exp1:

    if st.button("📊 Exportar estado seleccionado"):

        archivo_ppt = exportar_ppt_estado(
            tabla_base,
            fig_piezas,
            fig_montos,
            fig_ordenes,
            fig_claves,
            estado_sel
        )

        st.success("PPT del estado generado correctamente ✅")

        with open(archivo_ppt, "rb") as f:
            st.download_button(
                "⬇ Descargar PPT del estado",
                f,
                file_name=archivo_ppt
            )

with col_exp2:

    if st.button("🌎 Exportar TODOS los estados"):

        archivo_total = exportar_ppt_todos_estados()

        st.success("PowerPoint nacional generado correctamente ✅")

        with open(archivo_total, "rb") as f:
            st.download_button(
                "⬇ Descargar PowerPoint Nacional",
                f,
                file_name=archivo_total
            )